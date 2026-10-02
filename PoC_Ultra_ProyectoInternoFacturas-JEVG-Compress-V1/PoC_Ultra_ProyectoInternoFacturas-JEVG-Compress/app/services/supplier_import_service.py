"""Carga masiva del catalogo de proveedores desde la plantilla de Excel (HU-01).

Tres fases: (1) archivo: extension, tamano, paquete ZIP, XML, plantilla y limites; (2) filas: normalizacion y
reglas por columna; (3) duplicados en el archivo y contra el catalogo. En modo strict nada se escribe si hay filas
con errores; en modo partial, que el Administrador confirma para el mismo archivo (verificado por SHA-256), se
registran solo las filas validas. El archivo se procesa en memoria y no se conserva.
"""

from __future__ import annotations

import hashlib
import logging
import re
import time
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from datetime import time as clock_time
from io import BytesIO
from pathlib import PurePath

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from sqlalchemy import func, select, tuple_
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.constants import FOREIGN_TAX_ID_FORMAT, PHONE_FORMAT, SupplierOrigin, SupplierStatus, SupplierType
from app.core.countries import COUNTRIES
from app.models import Supplier, User
from app.services.audit_service import audit
from app.services.supplier_template import HEADERS, MAX_FILE_MB, MAX_ROWS, SHEET_NAME

logger = logging.getLogger(__name__)

STRICT, PARTIAL = "strict", "partial"
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ERRORS_SHOWN = 200

MSG_EXTENSION = "Solo se aceptan archivos de Excel (.xlsx) generados con la plantilla"
MSG_TOO_LARGE = f"El archivo excede {MAX_FILE_MB} MB"
MSG_NOT_WORKBOOK = "El archivo no es un libro de Excel válido"
MSG_UNCOMPRESSED = "El contenido del archivo excede el tamaño permitido"
MSG_TEMPLATE = "El archivo no corresponde a la plantilla vigente. Descargue la plantilla e intente de nuevo."
MSG_EMPTY = "El archivo no contiene proveedores"
MSG_TOO_MANY = f"El archivo excede el máximo de {MAX_ROWS} proveedores por carga"
MSG_MODE = "Modo de carga no válido"
MSG_ROW_ERRORS = "Algunas filas contienen errores o no se han podido leer correctamente."
MSG_CHANGED = "El archivo cambió desde la validación. Vuelva a cargarlo."
MSG_CONFLICT = "Otro proceso registró proveedores de este archivo mientras se procesaba. Vuelva a cargarlo."

ORIGINS = {"nacional": SupplierOrigin.NATIONAL, "internacional": SupplierOrigin.INTERNATIONAL}
PERSON_TYPES = {"fisica": SupplierType.PERSONA_FISICA, "moral": SupplierType.PERSONA_MORAL}
AGREEMENT = {"si": True, "no": False, "": False}
RFC_PATTERN = re.compile(r"^[A-ZÑ&]{3,4}[0-9]{6}[A-Z0-9]{3}$")
GENERIC_RFCS = {"XAXX010101000", "XEXX010101000"}
RFC_LENGTH = {SupplierType.PERSONA_MORAL: (12, "moral"), SupplierType.PERSONA_FISICA: (13, "física")}
TAX_ID_PATTERN = re.compile(FOREIGN_TAX_ID_FORMAT)
PHONE_PATTERN = re.compile(PHONE_FORMAT)
SPACES = re.compile(r"\s+")
XML_DECLARATIONS = re.compile(rb"<!(doctype|entity)", re.IGNORECASE)
COLUMN_ORDER = {header: index for index, header in enumerate(HEADERS)}


class ImportFileError(Exception):
    """El archivo completo se rechaza (HTTP 400) sin validar filas ni registrar proveedores."""

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


class Formula:
    """Celda con formula: se rechaza sin evaluarla."""


class DateValue:
    """Celda con fecha: Excel la convirtio; debe capturarse como texto."""


@dataclass
class RowError:
    row: int
    column: str
    message: str

    def as_dict(self) -> dict:
        return {
            "row": self.row,
            "column": self.column,
            "message": self.message,
            "text": f"Fila {self.row} · {self.column}: {self.message}",
        }


@dataclass
class Row:
    number: int
    raw: dict[str, object]
    values: dict = field(default_factory=dict)
    errors: list[RowError] = field(default_factory=list)
    skipped: bool = False

    def error(self, column: str, message: str) -> None:
        self.errors.append(RowError(self.number, column, message))

    @property
    def identifier(self) -> str:
        values = self.values
        return values.get("rfc") or f"{values.get('country')} {values.get('foreign_tax_id')}"


# --- Fase 1: archivo ---------------------------------------------------------------------------------------------


def read_rows(filename: str | None, content: bytes) -> list[Row]:
    if PurePath(filename or "").suffix.lower() != ".xlsx":
        raise ImportFileError(MSG_EXTENSION)
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(MSG_TOO_LARGE)
    _inspect_package(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # cualquier fallo del lector significa "no es un libro valido"
        _log_read_failure(exc)
        raise ImportFileError(MSG_NOT_WORKBOOK) from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ImportFileError(MSG_TEMPLATE)
        sheet = workbook[SHEET_NAME]
        _check_headers(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        rows = _data_rows(sheet)
    except ImportFileError:
        raise
    except Exception as exc:
        _log_read_failure(exc)
        raise ImportFileError(MSG_NOT_WORKBOOK) from exc
    finally:
        workbook.close()
    if not rows:
        raise ImportFileError(MSG_EMPTY)
    return rows


def _inspect_package(content: bytes, log_failure=None) -> None:
    """Revisa el ZIP antes de abrirlo: tamano descomprimido acotado y XML sin DTD ni entidades (Excel no los usa).
    `log_failure` permite a otra carga (catalogos, HU-07) registrar su propio evento de log."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as package:
            members = package.infolist()
            if "xl/workbook.xml" not in {member.filename for member in members}:
                raise ImportFileError(MSG_NOT_WORKBOOK)
            # ZipExtFile nunca entrega mas bytes que el tamano declarado: la suma acota toda la lectura.
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise ImportFileError(MSG_UNCOMPRESSED)
            for member in members:
                if member.filename.endswith((".xml", ".rels")) and XML_DECLARATIONS.search(package.read(member)):
                    raise ImportFileError(MSG_NOT_WORKBOOK)
    except ImportFileError:
        raise
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, EOFError, RuntimeError, ValueError) as exc:
        (log_failure or _log_read_failure)(exc)
        raise ImportFileError(MSG_NOT_WORKBOOK) from exc


def _check_headers(header_row: tuple) -> None:
    found = [str(value).strip() if value is not None else "" for value in header_row]
    while found and not found[-1]:
        found.pop()
    if [value.casefold() for value in found] != [header.casefold() for header in HEADERS]:
        raise ImportFileError(MSG_TEMPLATE)


def _data_rows(sheet) -> list[Row]:
    rows = []
    # max_col acota cada fila a las columnas de la plantilla aunque la hoja declare una dimension enorme.
    for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(HEADERS)), start=2):
        raw = {header: _cell_value(cell) for header, cell in zip(HEADERS, cells, strict=False)}
        if all(value is None for value in raw.values()):
            continue
        if len(rows) == MAX_ROWS:
            raise ImportFileError(MSG_TOO_MANY)
        rows.append(Row(number, {header: raw.get(header) for header in HEADERS}))
    return rows


def _cell_value(cell) -> object:
    value = getattr(cell, "value", None)
    if getattr(cell, "data_type", None) == "f":
        return Formula
    if isinstance(value, (datetime, date, clock_time, timedelta)):
        return DateValue
    if isinstance(value, str) and not value.strip():
        return None
    return value


def _log_read_failure(exc: Exception) -> None:
    logger.warning(
        "supplier_import.read_failed",
        extra={"event": "supplier_import.read_failed", "error_type": type(exc).__name__, "error": str(exc)},
    )


# --- Fase 2: filas -----------------------------------------------------------------------------------------------


def validate_rows(rows: list[Row]) -> None:
    for row in rows:
        _validate_row(row)
    _mark_duplicates(rows)
    for row in rows:
        row.errors.sort(key=lambda error: COLUMN_ORDER[error.column])


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return SPACES.sub(" ", str(value)).strip()


def _plain(text: str) -> str:
    """Minusculas y sin acentos, para comparar valores de lista."""
    return "".join(c for c in unicodedata.normalize("NFKD", text.casefold()) if not unicodedata.combining(c))


def _validate_row(row: Row) -> None:
    texts: dict[str, str | None] = {}
    for header, value in row.raw.items():
        if value is Formula:
            row.error(header, "la celda contiene una fórmula; capture el valor")
            texts[header] = None
        elif value is DateValue:
            row.error(header, "la celda contiene una fecha; capture el valor como texto")
            texts[header] = None
        else:
            texts[header] = _text(value)

    origin = _choice(row, texts, "Origen", ORIGINS, "use Nacional o Internacional")
    person_type = _choice(row, texts, "Tipo de persona", PERSON_TYPES, "use Física o Moral")
    agreement = _choice(row, texts, "Convenio de confidencialidad", AGREEMENT, "use Sí o No", required=False)
    business_name = _required(row, texts, "Razón social")
    if business_name is not None and len(business_name) < 2:
        row.error("Razón social", "es demasiado corto")
    elif business_name is not None and len(business_name) > 250:
        row.error("Razón social", "es demasiado largo")
    email = _email(row, texts)
    phone = texts["Teléfono"]
    if phone and not PHONE_PATTERN.fullmatch(phone):
        row.error("Teléfono", "use de 7 a 30 caracteres: dígitos, espacios, +, (, ) o -")
    notes = texts["Notas"]
    if notes and len(notes) > 1000:
        row.error("Notas", "admite a lo sumo 1000 caracteres")

    rfc, tax_id, country = (_upper(texts, header) for header in ("RFC", "Identificador fiscal extranjero", "País"))
    if origin == SupplierOrigin.NATIONAL:
        _national_identity(row, rfc, tax_id, country, person_type)
        country = "MX"
    elif origin == SupplierOrigin.INTERNATIONAL:
        _international_identity(row, rfc, tax_id, country)

    row.values = {
        "origin": origin,
        "supplier_type": person_type,
        "business_name": business_name,
        "rfc": rfc or None,
        "foreign_tax_id": tax_id or None,
        "country": country or None,
        "email": email,
        "phone": phone or None,
        "confidentiality_agreement": bool(agreement),
        "notes": notes or None,
    }


def _upper(texts: dict, header: str) -> str | None:
    """None si la celda ya se rechazo (formula o fecha); "" si esta vacia."""
    text = texts[header]
    return text.upper() if text is not None else None


def _required(row: Row, texts: dict, header: str) -> str | None:
    text = texts[header]
    if text == "":
        row.error(header, "es obligatorio")
    return text or None


def _choice(row: Row, texts: dict, header: str, options: dict, message: str, required: bool = True):
    text = texts[header]
    if text is None:
        return None
    if not text and required:
        row.error(header, "es obligatorio")
        return None
    if _plain(text) not in options:
        row.error(header, message)
        return None
    return options[_plain(text)]


def _email(row: Row, texts: dict) -> str | None:
    email = _required(row, texts, "Correo electrónico")
    if email is None:
        return None
    email = email.lower()
    if len(email) > 255:
        row.error("Correo electrónico", "es demasiado largo")
        return None
    try:
        validate_email(email, check_deliverability=False)
    except EmailNotValidError:
        row.error("Correo electrónico", "no es un correo válido")
        return None
    return email


def _national_identity(row: Row, rfc, tax_id, country, person_type) -> None:
    if rfc == "":
        row.error("RFC", "es obligatorio para proveedores nacionales")
    elif rfc in GENERIC_RFCS:
        row.error("RFC", "no se permite un RFC genérico; un proveedor extranjero se registra con Origen Internacional")
    elif rfc is not None and not (RFC_PATTERN.fullmatch(rfc) and _valid_rfc_date(rfc[-9:-3])):
        row.error("RFC", "tiene un formato inválido")
    elif rfc is not None and person_type in RFC_LENGTH and len(rfc) != RFC_LENGTH[person_type][0]:
        length, label = RFC_LENGTH[person_type]
        row.error("RFC", f"una persona {label} debe tener un RFC de {length} caracteres")
    if tax_id:
        row.error("Identificador fiscal extranjero", "debe quedar vacío para proveedores nacionales")
    if country and country != "MX":
        row.error("País", "debe quedar vacío o ser MX para proveedores nacionales")


def _valid_rfc_date(digits: str) -> bool:
    year, month, day = int(digits[:2]), int(digits[2:4]), int(digits[4:])
    for century in (1900, 2000):
        try:
            date(century + year, month, day)
            return True
        except ValueError:
            continue
    return False


def _international_identity(row: Row, rfc, tax_id, country) -> None:
    if tax_id == "":
        row.error("Identificador fiscal extranjero", "es obligatorio para proveedores internacionales")
    elif tax_id is not None and not TAX_ID_PATTERN.fullmatch(tax_id):
        row.error(
            "Identificador fiscal extranjero",
            "use hasta 40 caracteres: letras, dígitos, espacios, puntos, guiones o diagonales",
        )
    if country == "":
        row.error("País", "es obligatorio para proveedores internacionales")
    elif country == "MX":
        row.error("País", "un proveedor internacional no puede tener país MX")
    elif country is not None and country not in COUNTRIES:
        row.error("País", "use el código ISO de dos letras (p. ej. US)")
    if rfc:
        row.error("RFC", "debe quedar vacío para proveedores internacionales")


def _mark_duplicates(rows: list[Row]) -> None:
    """Un valor repetido dentro del archivo invalida todas sus apariciones (RD-10)."""
    keys: dict[tuple, list[Row]] = {}
    for row in rows:
        values = row.values
        if values["rfc"]:
            keys.setdefault(("RFC", values["rfc"]), []).append(row)
        if values["origin"] == SupplierOrigin.INTERNATIONAL and values["foreign_tax_id"] and values["country"]:
            keys.setdefault(("ID", values["country"], values["foreign_tax_id"]), []).append(row)
        if values["email"]:
            keys.setdefault(("EMAIL", values["email"]), []).append(row)
    for key, repeated in keys.items():
        if len(repeated) < 2:
            continue
        numbers = _join_numbers([row.number for row in repeated])
        column, message = {
            "RFC": ("RFC", f"repetido en las filas {numbers}"),
            "ID": ("Identificador fiscal extranjero", f"repetido para el mismo país en las filas {numbers}"),
            "EMAIL": ("Correo electrónico", f"repetido en las filas {numbers}"),
        }[key[0]]
        for row in repeated:
            row.error(column, message)


def _join_numbers(numbers: list[int], limit: int = 5) -> str:
    shown = [str(number) for number in numbers[:limit]]
    if len(numbers) > limit:
        return f"{', '.join(shown)} y {len(numbers) - limit} más"
    return f"{', '.join(shown[:-1])} y {shown[-1]}"


# --- Fase 3: catalogo --------------------------------------------------------------------------------------------


def classify(db: Session, rows: list[Row]) -> None:
    """Filas sin errores: omitidas si ya existen; error si su correo ya esta en uso; nuevas en otro caso."""
    clean = [row for row in rows if not row.errors]
    rfcs = {row.values["rfc"] for row in clean if row.values["rfc"]}
    pairs = {(row.values["country"], row.values["foreign_tax_id"]) for row in clean if row.values["foreign_tax_id"]}
    existing_rfcs = set(db.scalars(select(Supplier.rfc).where(Supplier.rfc.in_(rfcs)))) if rfcs else set()
    existing_pairs = (
        set(
            db.execute(
                select(Supplier.country, Supplier.foreign_tax_id).where(
                    tuple_(Supplier.country, Supplier.foreign_tax_id).in_(pairs)
                )
            ).tuples()
        )
        if pairs
        else set()
    )
    for row in clean:
        values = row.values
        row.skipped = values["rfc"] in existing_rfcs or (values["country"], values["foreign_tax_id"]) in existing_pairs
    emails = {row.values["email"] for row in clean if not row.skipped}
    used = set()
    if emails:
        used |= set(db.scalars(select(func.lower(Supplier.email)).where(func.lower(Supplier.email).in_(emails))))
        used |= set(db.scalars(select(func.lower(User.email)).where(func.lower(User.email).in_(emails))))
    for row in clean:
        if not row.skipped and row.values["email"] in used:
            row.error("Correo electrónico", "ya está registrado para otro proveedor o usuario")


def register(db: Session, rows: list[Row], user, sha256: str, summary: dict) -> None:
    suppliers = [Supplier(**row.values, status=SupplierStatus.REGISTERED) for row in rows]
    db.add_all(suppliers)
    db.flush()
    for supplier in suppliers:
        audit(
            db,
            "SUPPLIER_CREATED",
            "Supplier",
            supplier.id,
            user.id,
            new={"source": "bulk_import", "import_sha256": sha256},
        )
    audit(db, "SUPPLIER_BULK_IMPORTED", "SupplierImport", sha256, user.id, new=summary)
    db.commit()


# --- Orquestacion ------------------------------------------------------------------------------------------------


def import_suppliers(
    db: Session, user, filename: str | None, content: bytes, mode: str, expected_sha256: str | None
) -> tuple[int, dict]:
    """Devuelve (codigo HTTP, cuerpo JSON) segun el contrato D12 del diseno."""
    started = time.monotonic()
    sha256 = hashlib.sha256(content).hexdigest()
    event = {"mode": mode, "rows": 0, "registered": 0, "skipped": 0, "invalid": 0, "size_bytes": len(content)}
    if mode not in (STRICT, PARTIAL):
        return _finish(400, {"detail": MSG_MODE}, event, "rejected", started)
    if mode == PARTIAL and expected_sha256 != sha256:
        return _finish(409, {"detail": MSG_CHANGED}, event, "rejected", started)
    try:
        rows = read_rows(filename, content)
    except ImportFileError as exc:
        return _finish(400, {"detail": exc.message}, event, "rejected", started)
    validate_rows(rows)
    classify(db, rows)
    invalid = [row for row in rows if row.errors]
    new = [row for row in rows if not row.errors and not row.skipped]
    skipped = [
        {"row": row.number, "identifier": row.identifier, "business_name": row.values["business_name"]}
        for row in rows
        if row.skipped
    ]
    errors = sorted((error for row in invalid for error in row.errors), key=lambda e: (e.row, COLUMN_ORDER[e.column]))
    shown = [error.as_dict() for error in errors[:MAX_ERRORS_SHOWN]]
    hidden = max(0, len(errors) - MAX_ERRORS_SHOWN)
    event.update(rows=len(rows), skipped=len(skipped), invalid=len(invalid))
    if invalid and mode == STRICT:
        body = {
            "detail": MSG_ROW_ERRORS,
            "rows": len(rows),
            "registrable": len(new),
            "invalid": len(invalid),
            "errors": shown,
            "hidden_errors": hidden,
            "skipped": skipped,
            "sha256": sha256,
        }
        return _finish(400, body, event, "rejected", started)
    event["registered"] = len(new)
    audit_summary = {
        "mode": mode,
        "sha256": sha256,
        "size_bytes": len(content),
        "rows": len(rows),
        "created": len(new),
        "skipped": len(skipped),
        "invalid": len(invalid),
    }
    try:
        register(db, new, user, sha256, audit_summary)
    except IntegrityError:
        db.rollback()
        event["registered"] = 0
        return _finish(409, {"detail": MSG_CONFLICT}, event, "conflict", started)
    body = {
        "result": "partial" if invalid else "imported",
        "mode": mode,
        "rows": len(rows),
        "created": len(new),
        "skipped": skipped,
        "invalid": len(invalid),
        "errors": shown,
        "hidden_errors": hidden,
        "summary": summary_text(len(rows), len(new), len(skipped), len(invalid)),
    }
    return _finish(200, body, event, body["result"], started)


def summary_text(rows: int, created: int, skipped: int, invalid: int) -> str:
    parts = [
        _count(rows, "fila leída", "filas leídas"),
        _count(created, "proveedor registrado", "proveedores registrados"),
        _count(skipped, "omitido", "omitidos"),
    ]
    if invalid:
        parts.append(f"{invalid} con errores")
    return " · ".join(parts)


def _count(number: int, singular: str, plural: str) -> str:
    return f"{number} {singular if number == 1 else plural}"


def _finish(status_code: int, body: dict, event: dict, result: str, started: float) -> tuple[int, dict]:
    """Evento supplier.bulk_import sin datos del archivo: solo modo, resultado, conteos, tamano y duracion.

    El conteo de altas se llama "registered": "created" es un atributo reservado de LogRecord.
    """
    duration_ms = round((time.monotonic() - started) * 1000)
    logger.info(
        "supplier.bulk_import",
        extra={"event": "supplier.bulk_import", **event, "result": result, "duration_ms": duration_ms},
    )
    return status_code, body
