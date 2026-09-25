"""Formato predefinido de la carga masiva de proveedores (HU-01): encabezados, limites y plantilla de Excel.

El sistema genera la plantilla en cada descarga; no se versiona un binario. La hoja Proveedores lleva solo los
encabezados: los ejemplos van en la hoja Instrucciones, que nunca se importa.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.countries import COUNTRIES

TEMPLATE_VERSION = "v1"
TEMPLATE_FILENAME = f"plantilla_carga_proveedores_{TEMPLATE_VERSION}.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAME = "Proveedores"
INSTRUCTIONS_SHEET = "Instrucciones"
MAX_FILE_MB = 5
MAX_ROWS = 1000

HEADERS = (
    "Origen",
    "Tipo de persona",
    "Razón social",
    "RFC",
    "Identificador fiscal extranjero",
    "País",
    "Correo electrónico",
    "Teléfono",
    "Convenio de confidencialidad",
    "Notas",
)
# Formato de texto para que Excel no convierta ni recorte los valores (ceros iniciales, numeros largos).
TEXT_COLUMNS = ("RFC", "Identificador fiscal extranjero", "Teléfono")
LIST_VALUES = {
    "Origen": ("Nacional", "Internacional"),
    "Tipo de persona": ("Física", "Moral"),
    "Convenio de confidencialidad": ("Sí", "No"),
}
COLUMN_HELP = (
    ("Sí", "Nacional o Internacional."),
    ("Sí", "Física o Moral."),
    ("Sí", "Nombre o razón social, de 2 a 250 caracteres."),
    ("Nacional: sí. Internacional: vacío", "12 caracteres (persona moral) o 13 (persona física). Sin RFC genéricos."),
    (
        "Internacional: sí. Nacional: vacío",
        "Identificador fiscal del país de origen: hasta 40 letras, dígitos, espacios, puntos, guiones o diagonales.",
    ),
    (
        "Internacional: sí. Nacional: vacío o MX",
        "Código ISO de dos letras (lista abajo). MX no aplica a Internacional.",
    ),
    ("Sí", "Correo del proveedor. No puede usarlo otro proveedor ni otro usuario del portal."),
    ("No", "De 7 a 30 caracteres: dígitos, espacios, +, (, ) o -."),
    ("No", "Sí o No. Vacío equivale a No."),
    ("No", "Hasta 1000 caracteres."),
)
EXAMPLES = (
    (
        "Nacional",
        "Moral",
        "Servicios Digitales del Norte SA de CV",
        "SDN200315AB1",
        "",
        "",
        "contacto@sdn.mx",
        "55 1234 5678",
        "Sí",
        "",
    ),
    (
        "Internacional",
        "Moral",
        "Northwind Consulting LLC",
        "",
        "12-3456789",
        "US",
        "billing@northwind.example",
        "+1 206 555 0100",
        "No",
        "",
    ),
)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="12243A")


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        header, letter = cell.value, cell.column_letter
        sheet.column_dimensions[letter].width = max(16, len(header) + 4)
        if header in TEXT_COLUMNS:
            sheet.column_dimensions[letter].number_format = "@"
        if header in LIST_VALUES:
            options = LIST_VALUES[header]
            validation = DataValidation(
                type="list",
                formula1=f'"{",".join(options)}"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle=header,
                error=f"Use {' o '.join(options)}.",
            )
            validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
            sheet.add_data_validation(validation)
    sheet.freeze_panes = "A2"
    _instructions(workbook.create_sheet(INSTRUCTIONS_SHEET))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _instructions(sheet) -> None:
    bold = Font(bold=True)
    sheet.append([f"Plantilla de carga masiva de proveedores · versión {TEMPLATE_VERSION}"])
    sheet.append([f"Capture los proveedores en la hoja {SHEET_NAME} a partir de la fila 2."])
    sheet.append(["No cambie los encabezados, su orden ni el nombre de la hoja."])
    sheet.append(
        [
            f"Máximo {MAX_ROWS} proveedores y {MAX_FILE_MB} MB por archivo. "
            "Los proveedores que ya existen en el catálogo se omiten sin modificarse."
        ]
    )
    sheet.append([])
    sheet.append(["Columna", "Obligatoria", "Descripción"])
    for header, (required, description) in zip(HEADERS, COLUMN_HELP, strict=True):
        sheet.append([header, required, description])
    sheet.append([])
    sheet.append(["Ejemplos (esta hoja no se importa)"])
    sheet.append(list(HEADERS))
    for example in EXAMPLES:
        sheet.append(list(example))
    sheet.append([])
    sheet.append(["Código de país", "País"])
    for code, name in COUNTRIES.items():
        sheet.append([code, name])
    for row in sheet.iter_rows():
        if row[0].value in ("Columna", "Código de país") or (row[0].value or "").startswith(("Plantilla", "Ejemplos")):
            for cell in row:
                cell.font = bold
    for letter, width in (("A", 32), ("B", 38), ("C", 90)):
        sheet.column_dimensions[letter].width = width
