"""Catalogos de referencia (HU-07): monedas, usos de CFDI, formas y metodos de pago y regimenes fiscales.

- Alta, edicion de la descripcion y cambio de estado de claves, que nunca se borran. Una clave que usan las Reglas de
  Validacion (HU-06) no se puede desactivar, ni la ultima moneda activa (D4).
- Carga desde Excel todo o nada (D7): agrega claves nuevas y actualiza descripcion y estado de las existentes. Reutiliza
  la lectura segura de la carga masiva de proveedores (HU-01).

Las escrituras toman un bloqueo consultivo de transaccion. Estas funciones no confirman la transaccion: lo hace el
router, salvo la carga, que confirma y despues registra su evento de log.
"""

import logging
import re
import time
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePath

from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.core.constants import CATALOG_CODE_FORMATS, CATALOG_LABELS, CatalogType
from app.core.errors import BusinessRuleError, InvalidInputError, NotFoundError
from app.models import CatalogEntry, ValidationRule, now_utc
from app.repositories.pagination import Page, paginate, search
from app.rules.definitions import DEFINITIONS
from app.services.audit_service import audit

# Lectura segura del .xlsx compartida con la carga masiva de proveedores (HU-01): inspeccion del paquete ZIP
# (tamano descomprimido, sin DTD ni entidades), celdas con formula o fecha y el error de archivo completo.
from app.services.supplier_import_service import (
    MAX_FILE_BYTES,
    MSG_EXTENSION,
    MSG_NOT_WORKBOOK,
    MSG_TOO_LARGE,
    DateValue,
    Formula,
    ImportFileError,
    _cell_value,
    _inspect_package,
    _plain,
    _text,
)

logger = logging.getLogger(__name__)

# Clave del bloqueo consultivo que serializa las escrituras de los catalogos (pg_advisory_xact_lock).
CONFIG_LOCK_KEY = 7_0700_0001
ENTITY = "CatalogEntry"
NAME_MAX_LENGTH = 150
SHEET_NAME = "Catalogo"
HEADERS = ("Clave", "Descripción", "Activo")
MAX_ROWS = 1000
MAX_ERRORS_SHOWN = 200
ACTIVE_VALUES = {"si": True, "no": False}
# Catalogos numericos: Excel convierte "01" en 1; la clave se completa con ceros a su longitud.
NUMERIC_WIDTH = {CatalogType.PAYMENT_FORM: 2, CatalogType.TAX_REGIME: 3}

MSG_NOT_FOUND = "El catálogo no existe."
MSG_ENTRY_NOT_FOUND = "La clave no existe en este catálogo."
MSG_TEMPLATE = "El archivo no corresponde a la plantilla del catálogo. Descargue la plantilla e intente de nuevo."
MSG_EMPTY = "El archivo no contiene claves"
MSG_TOO_MANY = f"El archivo excede el máximo de {MAX_ROWS} claves por carga"
MSG_LAST_CURRENCY = "El catálogo de monedas debe conservar al menos una moneda activa."
# La moneda de un contrato o de una factura es siempre una clave activa del catalogo (ajustes-finales-configuracion).
MSG_INACTIVE_CURRENCY = "Moneda: la clave no está activa en el catálogo"


def catalog_for_code(code: str) -> CatalogType:
    """Catalogo de la ruta; uno desconocido es un 404 (no un 422 de FastAPI)."""
    try:
        return CatalogType(code)
    except ValueError:
        raise NotFoundError(MSG_NOT_FOUND) from None


# --- Lectura ------------------------------------------------------------------------------------------------------


def entries(db: Session, catalog: CatalogType) -> list[CatalogEntry]:
    return list(db.scalars(select(CatalogEntry).where(CatalogEntry.catalog == catalog).order_by(CatalogEntry.code)))


def page_entries(db: Session, catalog: CatalogType, q: str = "", page: int = 1) -> Page[CatalogEntry]:
    """Claves de la pantalla del catalogo: busqueda por clave o descripcion y paginacion en SQL. La plantilla de
    Excel sigue usando `entries` (todas las claves)."""
    stmt = select(CatalogEntry).where(CatalogEntry.catalog == catalog).order_by(CatalogEntry.code, CatalogEntry.id)
    return paginate(db, search(stmt, q, CatalogEntry.code, CatalogEntry.name), page)


@dataclass(frozen=True)
class CatalogCounts:
    total: int
    active: int


def counts(db: Session, catalog: CatalogType) -> CatalogCounts:
    """Claves del catalogo y cuantas estan activas, sobre todo el catalogo (no sobre la pagina)."""
    total, active = db.execute(
        select(func.count(), func.count().filter(CatalogEntry.is_active)).where(CatalogEntry.catalog == catalog)
    ).one()
    return CatalogCounts(total, active)


def active_codes(db: Session, catalog: CatalogType) -> set[str]:
    stmt = select(CatalogEntry.code).where(CatalogEntry.catalog == catalog, CatalogEntry.is_active.is_(True))
    return set(db.scalars(stmt))


def active_entries(db: Session, catalog: CatalogType) -> list[CatalogEntry]:
    return [entry for entry in entries(db, catalog) if entry.is_active]


@dataclass(frozen=True)
class CatalogSummary:
    catalog: CatalogType
    label: str
    active: int
    total: int


def summaries(db: Session) -> list[CatalogSummary]:
    counts = {
        catalog: (active or 0, total)
        for catalog, active, total in db.execute(
            select(
                CatalogEntry.catalog,
                func.count().filter(CatalogEntry.is_active.is_(True)),
                func.count(),
            ).group_by(CatalogEntry.catalog)
        )
    }
    return [CatalogSummary(c, CATALOG_LABELS[c], *counts.get(c, (0, 0))) for c in CatalogType]


def in_use(db: Session) -> dict[CatalogType, set[str]]:
    """Claves que usa el valor esperado de una regla activa de Reglas de Validacion, de cualquier origen: no se pueden
    desactivar (D4). Una regla eliminada no protege sus claves. Una actividad economica si: el proveedor la conserva y
    el formulario de edicion la sigue ofreciendo."""
    used: dict[CatalogType, set[str]] = {catalog: set() for catalog in CatalogType}
    for rule in db.scalars(select(ValidationRule).where(ValidationRule.is_active.is_(True))):
        catalog = DEFINITIONS[rule.origin, rule.rule_code].catalog
        if catalog is not None:
            used[catalog].update(rule.parameter if isinstance(rule.parameter, list) else [rule.parameter])
    return used


# --- Validacion ---------------------------------------------------------------------------------------------------


def normalize_code(catalog: CatalogType, raw: object) -> tuple[str, str | None]:
    """Clave en mayusculas y su error de formato, si lo hay."""
    width = NUMERIC_WIDTH.get(catalog)
    if width and isinstance(raw, (int, float)) and not isinstance(raw, bool) and float(raw).is_integer():
        code = str(int(raw)).zfill(width)
    else:
        code = _text(raw).upper()
    if not code:
        return code, "es obligatoria"
    pattern, description = CATALOG_CODE_FORMATS[catalog]
    return code, None if re.fullmatch(pattern, code) else description


def normalize_name(raw: object) -> tuple[str, str | None]:
    name = _text(raw)
    if not name:
        return name, "es obligatoria"
    if len(name) > NAME_MAX_LENGTH:
        return name, f"admite hasta {NAME_MAX_LENGTH} caracteres"
    return name, None


def _lock(db: Session) -> None:
    db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": CONFIG_LOCK_KEY})


def _entry(db: Session, catalog: CatalogType, entry_id: int) -> CatalogEntry:
    entry = db.get(CatalogEntry, entry_id)
    if entry is None or entry.catalog != catalog:
        raise NotFoundError(MSG_ENTRY_NOT_FOUND)
    return entry


def _in_use_message(code: str) -> str:
    return f"La clave {code} está en uso en Reglas de Validación."


# --- Escritura individual -----------------------------------------------------------------------------------------


def create_entry(db: Session, catalog: CatalogType, raw_code: str, raw_name: str, user_id: int) -> CatalogEntry:
    code, code_error = normalize_code(catalog, raw_code)
    name, name_error = normalize_name(raw_name)
    errors = []
    if code_error:
        errors.append(f"Clave: {code_error}")
    if name_error:
        errors.append(f"Descripción: {name_error}")
    if errors:
        raise InvalidInputError(" ".join(errors))
    _lock(db)
    exists = db.scalar(select(CatalogEntry.id).where(CatalogEntry.catalog == catalog, CatalogEntry.code == code))
    if exists:
        raise BusinessRuleError(f"La clave {code} ya existe en este catálogo")
    entry = CatalogEntry(catalog=catalog, code=code, name=name, is_active=True, updated_by=user_id)
    db.add(entry)
    db.flush()
    audit(
        db,
        "CATALOG_ENTRY_CREATED",
        ENTITY,
        entry.id,
        user_id,
        new={"catalog": catalog.value, "code": code, "name": name},
    )
    return entry


def update_entry(db: Session, catalog: CatalogType, entry_id: int, raw_name: str, user_id: int) -> bool:
    """Cambia la descripcion; la clave no cambia. True si hubo cambios."""
    name, error = normalize_name(raw_name)
    if error:
        raise InvalidInputError(f"Descripción: {error}")
    _lock(db)
    entry = _entry(db, catalog, entry_id)
    if entry.name == name:
        return False
    audit(db, "CATALOG_ENTRY_UPDATED", ENTITY, entry.id, user_id, {"name": entry.name}, {"name": name})
    entry.name, entry.updated_by = name, user_id
    return True


def set_active(db: Session, catalog: CatalogType, entry_id: int, active: bool, user_id: int) -> bool:
    """Desactiva o reactiva una clave con un estado destino explicito. True si hubo cambios."""
    _lock(db)
    entry = _entry(db, catalog, entry_id)
    if entry.is_active == active:
        return False
    if not active:
        if entry.code in in_use(db)[catalog]:
            raise BusinessRuleError(_in_use_message(entry.code))
        if catalog == CatalogType.CURRENCY and active_codes(db, catalog) == {entry.code}:
            raise BusinessRuleError(MSG_LAST_CURRENCY)
    audit(
        db,
        "CATALOG_ENTRY_STATUS_CHANGED",
        ENTITY,
        entry.id,
        user_id,
        {"is_active": entry.is_active},
        {"is_active": active},
    )
    entry.is_active, entry.updated_by = active, user_id
    return True


# --- Carga desde Excel (D7) ---------------------------------------------------------------------------------------


@dataclass
class ImportRow:
    number: int
    code: str
    name: str
    active: bool
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportResult:
    rows: int
    created: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    unchanged: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    hidden_errors: int = 0

    @property
    def invalid(self) -> bool:
        return bool(self.errors)

    @property
    def summary(self) -> str:
        """Resumen de una carga aplicada, p. ej. "1 agregada, 1 actualizada, 0 sin cambios"."""
        created, updated = len(self.created), len(self.updated)
        return (
            f"{created} {'agregada' if created == 1 else 'agregadas'}, "
            f"{updated} {'actualizada' if updated == 1 else 'actualizadas'}, {len(self.unchanged)} sin cambios"
        )


def _log_read_failure(exc: Exception) -> None:
    logger.warning(
        "catalog_import.read_failed",
        extra={"event": "catalog_import.read_failed", "error_type": type(exc).__name__, "error": str(exc)},
    )


def read_rows(filename: str | None, content: bytes) -> list[tuple[int, tuple]]:
    """Filas de datos (numero, celdas) del archivo; ImportFileError si el archivo completo no es valido."""
    if PurePath(filename or "").suffix.lower() != ".xlsx":
        raise ImportFileError(MSG_EXTENSION)
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(MSG_TOO_LARGE)
    _inspect_package(content, log_failure=_log_read_failure)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # cualquier fallo del lector significa "no es un libro valido"
        _log_read_failure(exc)
        raise ImportFileError(MSG_NOT_WORKBOOK) from exc
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ImportFileError(MSG_TEMPLATE)
        sheet = workbook[SHEET_NAME]
        header = [_text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        while header and not header[-1]:
            header.pop()
        if [value.casefold() for value in header] != [name.casefold() for name in HEADERS]:
            raise ImportFileError(MSG_TEMPLATE)
        rows = []
        for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(HEADERS)), start=2):
            values = tuple(_cell_value(cell) for cell in cells) + (None,) * (len(HEADERS) - len(cells))
            if all(value is None for value in values):
                continue
            if len(rows) == MAX_ROWS:
                raise ImportFileError(MSG_TOO_MANY)
            rows.append((number, values))
    except ImportFileError:
        raise
    except Exception as exc:
        _log_read_failure(exc)
        raise ImportFileError(MSG_NOT_WORKBOOK) from exc
    finally:
        workbook.close()
    if not rows:
        raise ImportFileError(MSG_EMPTY)
    return rows


def _cell_error(value: object) -> str | None:
    if value is Formula:
        return "no admite fórmulas; capture el valor"
    if value is DateValue:
        return "Excel lo convirtió en fecha; captúrelo como texto"
    return None


def validate_rows(catalog: CatalogType, raw_rows: list[tuple[int, tuple]], used: set[str]) -> list[ImportRow]:
    rows, first_seen = [], {}
    for number, (raw_code, raw_name, raw_active) in raw_rows:
        errors = []
        code, code_error = normalize_code(catalog, None if _cell_error(raw_code) else raw_code)
        code_error = _cell_error(raw_code) or code_error
        name, name_error = normalize_name(None if _cell_error(raw_name) else raw_name)
        name_error = _cell_error(raw_name) or name_error
        flag = _plain(_text(raw_active)) or "si"
        active = ACTIVE_VALUES.get(flag)
        if code_error:
            errors.append(f"Fila {number} · Clave: {code_error}")
        elif code in first_seen:
            errors.append(
                f"Fila {number} · Clave: la clave {code} está repetida en el archivo (fila {first_seen[code]})"
            )
        else:
            first_seen[code] = number
        if name_error:
            errors.append(f"Fila {number} · Descripción: {name_error}")
        if active is None:
            errors.append(f"Fila {number} · Activo: use Sí o No")
        elif not active and code in used:
            errors.append(f"Fila {number} · Activo: la clave {code} está en uso en Reglas de Validación")
        rows.append(ImportRow(number, code, name, bool(active), errors))
    return rows


def import_catalog(
    db: Session, catalog: CatalogType, filename: str | None, content: bytes, user_id: int
) -> ImportResult:
    """Carga todo o nada. ImportFileError si el archivo no es valido; un resultado con `errors` si alguna fila no lo
    es (nada se aplica). Sin errores, aplica, audita, confirma y registra el evento de log."""
    started = time.perf_counter()
    try:
        raw_rows = read_rows(filename, content)
    except ImportFileError:
        _log_import(catalog, "rejected", ImportResult(0), 0, len(content), started)
        raise
    _lock(db)
    current = {entry.code: entry for entry in entries(db, catalog)}
    rows = validate_rows(catalog, raw_rows, in_use(db)[catalog])
    errors = [error for row in rows for error in row.errors]
    result = ImportResult(len(rows))
    if catalog == CatalogType.CURRENCY and not errors:
        final = {code for code, entry in current.items() if entry.is_active}
        final |= {row.code for row in rows if row.active}
        final -= {row.code for row in rows if not row.active}
        if not final:
            errors.append(MSG_LAST_CURRENCY)
    if errors:
        db.rollback()  # libera el bloqueo consultivo
        result.errors, result.hidden_errors = errors[:MAX_ERRORS_SHOWN], max(0, len(errors) - MAX_ERRORS_SHOWN)
        _log_import(catalog, "rejected", result, sum(bool(row.errors) for row in rows), len(content), started)
        return result
    now = now_utc()
    for row in rows:
        entry = current.get(row.code)
        if entry is None:
            db.add(
                CatalogEntry(catalog=catalog, code=row.code, name=row.name, is_active=row.active, updated_by=user_id)
            )
            result.created.append(row.code)
        elif (entry.name, entry.is_active) != (row.name, row.active):
            entry.name, entry.is_active, entry.updated_at, entry.updated_by = row.name, row.active, now, user_id
            result.updated.append(row.code)
        else:
            result.unchanged.append(row.code)
    if result.created or result.updated:
        audit(
            db,
            "CATALOG_IMPORTED",
            ENTITY,
            None,
            user_id,
            new={"catalog": catalog.value, "created": result.created, "updated": result.updated},
        )
    db.commit()
    _log_import(catalog, "imported", result, 0, len(content), started)
    return result


def _log_import(
    catalog: CatalogType, outcome: str, result: ImportResult, invalid: int, size: int, started: float
) -> None:
    logger.info(
        "catalog.import",
        extra={
            "event": "catalog.import",
            "catalog": catalog.value,
            "result": outcome,
            "rows": result.rows,
            "added": len(result.created),  # "created" es un atributo reservado de LogRecord
            "updated": len(result.updated),
            "unchanged": len(result.unchanged),
            "invalid": invalid,
            "size_bytes": size,
            "duration_ms": round((time.perf_counter() - started) * 1000),
        },
    )
