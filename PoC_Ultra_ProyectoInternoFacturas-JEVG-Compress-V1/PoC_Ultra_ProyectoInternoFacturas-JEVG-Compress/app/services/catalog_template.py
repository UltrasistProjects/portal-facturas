"""Plantilla de Excel de un catalogo de referencia (HU-07): la hoja "Catalogo" con las claves vigentes, lista para
editarse y volver a cargarse, y una hoja de instrucciones. Mismo estilo que la plantilla de proveedores (HU-01)."""

from collections.abc import Iterable
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.constants import CATALOG_CODE_FORMATS, CATALOG_LABELS, CatalogType
from app.models import CatalogEntry
from app.services.catalog_service import HEADERS, MAX_ROWS, SHEET_NAME
from app.services.supplier_template import HEADER_FILL, HEADER_FONT

INSTRUCTIONS_SHEET = "Instrucciones"
TEXT_FORMAT = "@"
WIDTHS = (12, 70, 10)


def template_filename(catalog: CatalogType) -> str:
    return f"catalogo_{catalog.value.lower()}.xlsx"


def build_catalog_template(catalog: CatalogType, entries: Iterable[CatalogEntry]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for entry in entries:
        sheet.append([entry.code, entry.name, "Sí" if entry.is_active else "No"])
    for index, width in enumerate(WIDTHS):
        sheet.column_dimensions["ABC"[index]].width = width
    # La clave se captura como texto: Excel convertiria "01" en 1. El formato de columna no crea filas vacias.
    sheet.column_dimensions["A"].number_format = TEXT_FORMAT
    for (cell,) in sheet.iter_rows(min_row=2, max_col=1):
        cell.number_format = TEXT_FORMAT
    flags = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    flags.add(f"C2:C{MAX_ROWS + 1}")
    sheet.add_data_validation(flags)
    sheet.freeze_panes = "A2"
    _instructions(workbook.create_sheet(INSTRUCTIONS_SHEET), catalog)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _instructions(sheet, catalog: CatalogType) -> None:
    lines = (
        f"Catálogo: {CATALOG_LABELS[catalog]}",
        "",
        f'1. Edite la hoja "{SHEET_NAME}". No cambie el nombre de la hoja ni los encabezados.',
        f"2. Clave: obligatoria; {CATALOG_CODE_FORMATS[catalog][1]}. No se repite en el archivo.",
        "3. Descripción: obligatoria, hasta 150 caracteres.",
        "4. Activo: Sí o No. Vacío equivale a Sí.",
        "5. Las claves nuevas se agregan; las existentes actualizan su descripción y su estado.",
        "6. Las claves que no vienen en el archivo no cambian: nunca se borran.",
        "7. Si una fila tiene errores no se aplica nada. Corrija y vuelva a cargar el archivo.",
        "8. Una clave en uso en Reglas de Validación no se puede desactivar.",
        f"Máximo {MAX_ROWS} filas y 5 MB por archivo.",
    )
    for line in lines:
        sheet.append([line])
    sheet.column_dimensions["A"].width = 100
