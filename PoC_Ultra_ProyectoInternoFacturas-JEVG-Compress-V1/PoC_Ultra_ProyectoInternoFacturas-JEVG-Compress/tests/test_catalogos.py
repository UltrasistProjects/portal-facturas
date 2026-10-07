"""Catalogos de referencia (HU-07, spec catalogos-referencia)."""

import html
import re
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.constants import CatalogType, SupplierOrigin
from app.core.database import SessionLocal
from app.models import AuditLog, CatalogEntry
from app.rules.xml_rules import xml_rules
from app.services import catalog_service as cs
from app.services.validation_rules_service import rule_set
from tests.conftest import csrf, login

pytestmark = pytest.mark.usefixtures("restore_validation_rules")

URL = "/admin/catalogs"


# --- Utilidades ---------------------------------------------------------------------------------------------------


def entry(catalog: str, code: str) -> CatalogEntry | None:
    with SessionLocal() as db:
        return db.scalar(select(CatalogEntry).where(CatalogEntry.catalog == catalog, CatalogEntry.code == code))


def audit_count() -> int:
    with SessionLocal() as db:
        return db.scalar(select(func.count(AuditLog.id)))


def last_audit(action: str) -> AuditLog:
    with SessionLocal() as db:
        return db.scalar(select(AuditLog).where(AuditLog.action == action).order_by(AuditLog.id.desc()))


def post(client, path: str, **data):
    return client.post(f"{URL}{path}", data={**data, "csrf_token": csrf(client, "/")}, follow_redirects=False)


def create(client, catalog: str, code: str, name: str):
    return post(client, f"/{catalog}", entry_code=code, name=name)


def set_status(client, catalog: str, code: str, active: bool):
    return post(client, f"/{catalog}/{entry(catalog, code).id}/status", active="true" if active else "false")


def currency_accepted(code: str) -> bool:
    with SessionLocal() as db:
        results = {r.rule_code: r for r in xml_rules({"currency": code}, None, rule_set(db, SupplierOrigin.NATIONAL))}
    return results["XML-007"].status == "PASS"


def workbook(rows, sheet: str = "Catalogo", headers=("Clave", "Descripción", "Activo")) -> bytes:
    book = Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def upload(client, catalog: str, content: bytes, filename: str = "catalogo.xlsx"):
    return client.post(
        f"{URL}/{catalog}/import",
        data={"csrf_token": csrf(client, "/")},
        files={"upload": (filename, content, "application/octet-stream")},
    )


# --- Acceso -------------------------------------------------------------------------------------------------------


def test_sin_acceso_para_pmo(client):
    login(client, "pmo@poc.local")
    token = csrf(client, "/")
    responses = [
        client.get(URL),
        client.get(f"{URL}/CURRENCY"),
        client.get(f"{URL}/CURRENCY/template"),
        client.post(f"{URL}/CURRENCY", data={"entry_code": "CAD", "name": "Dólar canadiense", "csrf_token": token}),
        client.post(f"{URL}/CURRENCY/{entry('CURRENCY', 'EUR').id}", data={"name": "X", "csrf_token": token}),
        client.post(
            f"{URL}/CURRENCY/{entry('CURRENCY', 'EUR').id}/status", data={"active": "false", "csrf_token": token}
        ),
        client.post(
            f"{URL}/CURRENCY/import",
            data={"csrf_token": token},
            files={"upload": ("c.xlsx", workbook([("CAD", "Dólar canadiense", "Sí")]), "application/octet-stream")},
        ),
    ]
    assert [r.status_code for r in responses] == [403] * 7
    assert entry("CURRENCY", "CAD") is None and entry("CURRENCY", "EUR").is_active


def test_alta_sin_token_csrf(client):
    login(client)
    response = client.post(f"{URL}/CURRENCY", data={"entry_code": "CAD", "name": "Dólar canadiense"})
    assert response.status_code == 403 and entry("CURRENCY", "CAD") is None


def test_catalogo_inexistente(client):
    login(client)
    assert client.get(f"{URL}/COUNTRY").status_code == 404
    assert create(client, "COUNTRY", "MX", "México").status_code == 404
    assert post(client, f"/CURRENCY/{entry('PAYMENT_METHOD', 'PUE').id}", name="Otro").status_code == 404


# --- Listado ------------------------------------------------------------------------------------------------------


def test_listado_de_catalogos(client):
    login(client)
    page = client.get(URL).text
    rows = re.findall(
        r"<strong>([^<]+)</strong><small class=\"mono\">(\w+)</small></td><td>(\d+)</td><td>(\d+)</td>", page
    )
    assert rows == [
        ("Monedas", "CURRENCY", "3", "3"),
        ("Usos de CFDI", "CFDI_USE", "24", "24"),
        ("Formas de pago", "PAYMENT_FORM", "22", "22"),
        ("Métodos de pago", "PAYMENT_METHOD", "2", "2"),
        ("Regímenes fiscales", "TAX_REGIME", "19", "19"),
        ("Actividades económicas", "INDUSTRY", "20", "20"),
    ]


def test_claves_en_uso(client):
    login(client)
    page = client.get(f"{URL}/PAYMENT_FORM").text
    row = re.search(r"<td class=\"mono\">99.*?</tr>", page, re.DOTALL).group(0)
    assert "En uso" in row and "disabled" in row
    other = re.search(r"<td class=\"mono\">03.*?</tr>", page, re.DOTALL).group(0)
    assert "En uso" not in other


# --- Alta, edicion y estado ---------------------------------------------------------------------------------------


def test_alta_de_una_moneda(client):
    login(client)
    assert not currency_accepted("CAD")
    response = create(client, "CURRENCY", " cad ", "Dólar canadiense")
    # El alta regresa buscando la clave creada (listados-paginados).
    assert response.status_code == 303 and response.headers["location"] == f"{URL}/CURRENCY?q=CAD&ok=created"
    created = entry("CURRENCY", "CAD")
    assert (created.name, created.is_active) == ("Dólar canadiense", True)
    assert currency_accepted("CAD")
    audit = last_audit("CATALOG_ENTRY_CREATED")
    assert audit.new_value == {"catalog": "CURRENCY", "code": "CAD", "name": "Dólar canadiense"}


@pytest.mark.parametrize(
    ("catalog", "code", "name", "message"),
    [
        ("PAYMENT_FORM", "9A", "Otra", "Clave: debe tener dos dígitos"),
        ("CFDI_USE", "G1", "Otro", "Clave: debe tener una o dos letras seguidas de dos dígitos"),
        ("CURRENCY", "CAD", "", "Descripción: es obligatoria"),
        ("TAX_REGIME", "", "x" * 151, "Clave: es obligatoria Descripción: admite hasta 150 caracteres"),
    ],
)
def test_alta_invalida(client, catalog, code, name, message):
    login(client)
    before = audit_count()
    response = create(client, catalog, code, name)
    assert response.status_code == 400 and message in html.unescape(response.text)
    assert audit_count() == before


def test_clave_repetida(client):
    login(client)
    response = create(client, "PAYMENT_METHOD", "pue", "Otro")
    assert response.status_code == 409 and "La clave PUE ya existe en este catálogo" in response.text


def test_edicion_de_la_descripcion(client):
    login(client)
    target = entry("CURRENCY", "USD")
    response = post(client, f"/CURRENCY/{target.id}", name="Dólar estadounidense")
    assert response.headers["location"] == f"{URL}/CURRENCY?ok=updated"
    assert entry("CURRENCY", "USD").name == "Dólar estadounidense"
    audit = last_audit("CATALOG_ENTRY_UPDATED")
    assert (audit.old_value, audit.new_value) == ({"name": "Dólar americano"}, {"name": "Dólar estadounidense"})
    before = audit_count()
    again = post(client, f"/CURRENCY/{target.id}", name="Dólar estadounidense")
    assert again.headers["location"] == f"{URL}/CURRENCY?ok=unchanged" and audit_count() == before


def test_clave_en_uso_no_se_desactiva(client):
    login(client)
    before = audit_count()
    response = set_status(client, "CFDI_USE", "G03", False)
    assert response.status_code == 409 and "La clave G03 está en uso en Reglas de Validación." in response.text
    assert entry("CFDI_USE", "G03").is_active and audit_count() == before


def test_desactivar_y_reactivar_una_moneda(client):
    login(client)
    assert set_status(client, "CURRENCY", "EUR", False).headers["location"] == f"{URL}/CURRENCY?ok=status"
    assert not entry("CURRENCY", "EUR").is_active and not currency_accepted("EUR")
    audit = last_audit("CATALOG_ENTRY_STATUS_CHANGED")
    assert (audit.old_value, audit.new_value) == ({"is_active": True}, {"is_active": False})
    assert set_status(client, "CURRENCY", "EUR", True).status_code == 303
    assert entry("CURRENCY", "EUR").is_active and currency_accepted("EUR")


def test_ultima_moneda_activa(client):
    login(client)
    set_status(client, "CURRENCY", "USD", False)
    set_status(client, "CURRENCY", "EUR", False)
    response = set_status(client, "CURRENCY", "MXN", False)
    assert response.status_code == 409 and cs.MSG_LAST_CURRENCY in response.text
    assert entry("CURRENCY", "MXN").is_active


def test_estado_invalido(client):
    login(client)
    response = post(client, f"/CURRENCY/{entry('CURRENCY', 'EUR').id}/status", active="quizas")
    assert response.status_code == 400 and "Estado inválido" in response.text


def test_clave_inactiva_no_se_ofrece_en_reglas(client):
    login(client)
    set_status(client, "PAYMENT_FORM", "01", False)
    page = client.get("/admin/rules").text
    assert '<option value="01"' not in page and '<option value="03"' in page


# --- Carga desde Excel --------------------------------------------------------------------------------------------


def test_plantilla_con_el_contenido_vigente(client):
    login(client)
    response = client.get(f"{URL}/PAYMENT_METHOD/template")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="catalogo_payment_method.xlsx"'
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Catalogo", "Instrucciones"]
    rows = list(book["Catalogo"].iter_rows(values_only=True))
    assert rows == [
        ("Clave", "Descripción", "Activo"),
        ("PPD", "Pago en parcialidades o diferido", "Sí"),
        ("PUE", "Pago en una sola exhibición", "Sí"),
    ]


def test_la_plantilla_se_carga_sin_cambios(client):
    login(client)
    content = client.get(f"{URL}/PAYMENT_FORM/template").content
    response = upload(client, "PAYMENT_FORM", content)
    assert response.status_code == 200 and "0 agregadas, 0 actualizadas, 22 sin cambios" in response.text


def test_carga_con_claves_nuevas_y_actualizadas(client):
    login(client)
    content = workbook([("USD", "Dólar estadounidense", "Sí"), ("jpy", "Yen japonés", None)])
    response = upload(client, "CURRENCY", content)
    assert response.status_code == 200 and "Carga aplicada: 1 agregada, 1 actualizada, 0 sin cambios." in response.text
    assert entry("CURRENCY", "JPY").is_active and entry("CURRENCY", "USD").name == "Dólar estadounidense"
    assert entry("CURRENCY", "MXN").name == "Peso mexicano" and entry("CURRENCY", "EUR").is_active
    audit = last_audit("CATALOG_IMPORTED")
    assert audit.new_value == {"catalog": "CURRENCY", "created": ["JPY"], "updated": ["USD"]}


def test_clave_numerica_se_completa_con_ceros(client):
    login(client)
    response = upload(client, "PAYMENT_FORM", workbook([(3, "Transferencia electrónica", "Sí")]))
    assert response.status_code == 200 and "1 actualizada" in response.text
    assert entry("PAYMENT_FORM", "03").name == "Transferencia electrónica" and entry("PAYMENT_FORM", "3") is None


def test_filas_con_errores_no_aplican_nada(client):
    login(client)
    before = audit_count()
    content = workbook(
        [
            ("9A", "Otra", "Sí"),
            ("99", "Por definir", "No"),
            ("01", "", "Sí"),
            ("02", "Cheque", "Quizás"),
            ("03", "Transferencia", "Sí"),
            ("03", "Transferencia", "Sí"),
            ("04", "=1+1", "Sí"),
        ]
    )
    response = upload(client, "PAYMENT_FORM", content)
    text = html.unescape(response.text)
    assert response.status_code == 400 and "No se aplicó la carga" in text
    for message in (
        "Fila 2 · Clave: debe tener dos dígitos",
        "Fila 3 · Activo: la clave 99 está en uso en Reglas de Validación",
        "Fila 4 · Descripción: es obligatoria",
        "Fila 5 · Activo: use Sí o No",
        "Fila 7 · Clave: la clave 03 está repetida en el archivo (fila 6)",
        "Fila 8 · Descripción: no admite fórmulas; capture el valor",
    ):
        assert message in text, message
    assert entry("PAYMENT_FORM", "99").is_active and entry("PAYMENT_FORM", "03").name.startswith("Transferencia elec")
    assert audit_count() == before


def test_carga_que_deja_monedas_sin_activas(client):
    login(client)
    content = workbook([("MXN", "Peso mexicano", "No"), ("USD", "Dólar", "No"), ("EUR", "Euro", "No")])
    response = upload(client, "CURRENCY", content)
    assert response.status_code == 400 and cs.MSG_LAST_CURRENCY in response.text
    assert entry("CURRENCY", "MXN").is_active


@pytest.mark.parametrize(
    ("content", "filename", "message"),
    [
        (workbook([("MXN", "Peso", "Sí")], sheet="Hoja1"), "c.xlsx", cs.MSG_TEMPLATE),
        (workbook([("MXN", "Peso", "Sí")], headers=("Codigo", "Nombre", "Activo")), "c.xlsx", cs.MSG_TEMPLATE),
        (workbook([]), "c.xlsx", cs.MSG_EMPTY),
        (workbook([("MXN", "Peso", "Sí")]), "c.xls", cs.MSG_EXTENSION),
        (b"no es un zip", "c.xlsx", cs.MSG_NOT_WORKBOOK),
        (b"x" * (cs.MAX_FILE_BYTES + 1), "c.xlsx", cs.MSG_TOO_LARGE),
    ],
    ids=["otra-hoja", "otros-encabezados", "sin-filas", "extension", "no-es-zip", "demasiado-grande"],
)
def test_archivo_invalido(client, content, filename, message):
    login(client)
    response = upload(client, "CURRENCY", content, filename)
    assert response.status_code == 400 and message in html.unescape(response.text)


def test_demasiadas_filas():
    rows = [(f"{index:03d}", "Régimen", "Sí") for index in range(cs.MAX_ROWS + 1)]
    with pytest.raises(cs.ImportFileError, match="máximo"):
        cs.read_rows("c.xlsx", workbook(rows))


def test_resumen_por_catalogo():
    with SessionLocal() as db:
        summaries = {s.catalog: (s.active, s.total) for s in cs.summaries(db)}
    assert summaries[CatalogType.PAYMENT_METHOD] == (2, 2)
