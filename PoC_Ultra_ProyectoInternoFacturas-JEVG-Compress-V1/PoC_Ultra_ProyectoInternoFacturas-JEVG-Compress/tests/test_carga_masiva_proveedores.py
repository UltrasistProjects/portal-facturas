"""Carga masiva de proveedores (HU-01, spec catalogo-proveedores). Los libros se generan en memoria y cada prueba usa
RFC, identificadores y correos unicos, porque la base de la sesion es compartida."""

import hashlib
import io
import json
import zipfile
from datetime import date
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.config import settings
from app.core.constants import SupplierOrigin, SupplierStatus, SupplierType
from app.core.database import SessionLocal
from app.models import AuditLog, Supplier, User
from app.services import supplier_import_service as service
from app.services.supplier_template import HEADERS, TEMPLATE_FILENAME, XLSX_MEDIA_TYPE
from tests.conftest import csrf, login

PAGE = "/suppliers/import"
LOG_FILE = settings.log_dir / "app.log"
LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
HOMOCLAVE = "ABCDEFGHJKLMNPQRSTUVWXYZ0123456789"
MSG_TEMPLATE = "El archivo no corresponde a la plantilla vigente. Descargue la plantilla e intente de nuevo."


# --- Utilidades --------------------------------------------------------------------------------------------------


def unique_rfc(moral: bool = True) -> str:
    token = uuid4().int
    letters = "".join(LETTERS[(token >> (5 * i)) % 26] for i in range(3 if moral else 4))
    homoclave = "".join(HOMOCLAVE[(token >> (40 + 6 * i)) % len(HOMOCLAVE)] for i in range(3))
    return f"{letters}200315{homoclave}"


def unique_email(domain: str = "proveedor.mx") -> str:
    return f"{uuid4().hex[:12]}@{domain}"


def national(**overrides) -> dict:
    values = {
        "Origen": "Nacional",
        "Tipo de persona": "Moral",
        "Razón social": "Proveedor de Prueba SA de CV",
        "RFC": unique_rfc(),
        "Correo electrónico": unique_email(),
    }
    return {**values, **overrides}


def international(**overrides) -> dict:
    values = {
        "Origen": "Internacional",
        "Tipo de persona": "Moral",
        "Razón social": "Foreign Consulting LLC",
        "Identificador fiscal extranjero": f"TX-{uuid4().hex[:10].upper()}",
        "País": "US",
        "Correo electrónico": unique_email("foreign.example"),
    }
    return {**values, **overrides}


def workbook(rows: list[dict | None], headers=HEADERS, sheet: str = "Proveedores") -> bytes:
    book = Workbook()
    worksheet = book.active
    worksheet.title = sheet
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append([None] * len(headers) if row is None else [row.get(header) for header in headers])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def admin(client):
    login(client)
    return client


def post_import(client, content: bytes, filename="proveedores.xlsx", mode=None, sha256=None, token=True):
    data = {"csrf_token": csrf(client, PAGE)} if token else {}
    if mode:
        data["mode"] = mode
    if sha256:
        data["expected_sha256"] = sha256
    files = {"upload": (filename, content, "application/octet-stream")}
    return client.post(PAGE, data=data, files=files, headers={"Accept": "application/json"})


def texts(response) -> list[str]:
    return [error["text"] for error in response.json()["errors"]]


def supplier_count() -> int:
    with SessionLocal() as db:
        return db.scalar(select(func.count()).select_from(Supplier))


def audit_count() -> int:
    with SessionLocal() as db:
        return db.scalar(select(func.count()).select_from(AuditLog))


def supplier_by(**criteria) -> Supplier | None:
    with SessionLocal() as db:
        return db.scalar(select(Supplier).filter_by(**criteria))


def existing_supplier(**overrides) -> Supplier:
    values = {
        "business_name": "Proveedor Existente SA de CV",
        "rfc": unique_rfc(),
        "supplier_type": SupplierType.PERSONA_MORAL,
        "email": unique_email(),
        **overrides,
    }
    with SessionLocal() as db:
        supplier = Supplier(**values)
        db.add(supplier)
        db.commit()
        return supplier


def log_offset() -> int:
    return LOG_FILE.stat().st_size if LOG_FILE.exists() else 0


def events_since(offset: int) -> list[dict]:
    with LOG_FILE.open(encoding="utf-8") as handle:
        handle.seek(offset)
        return [json.loads(line) for line in handle.read().splitlines() if line.strip()]


# --- Plantilla y acceso ------------------------------------------------------------------------------------------


def test_descarga_de_la_plantilla(admin):
    response = admin.get(f"{PAGE}/template")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert TEMPLATE_FILENAME in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Proveedores", "Instrucciones"]
    sheet = book["Proveedores"]
    assert list(sheet.iter_rows(values_only=True)) == [HEADERS]
    lists = {str(validation.sqref): validation.formula1 for validation in sheet.data_validations.dataValidation}
    assert lists == {"A2:A1001": '"Nacional,Internacional"', "B2:B1001": '"Física,Moral"', "I2:I1001": '"Sí,No"'}
    assert {sheet.column_dimensions[letter].number_format for letter in "DEH"} == {"@"}
    instructions = [row[0] for row in book["Instrucciones"].iter_rows(values_only=True)]
    assert any("versión v1" in (value or "") for value in instructions)
    assert "US" in instructions and "MX" in instructions


def test_los_ejemplos_no_se_importan(admin):
    before = supplier_count()
    response = post_import(admin, admin.get(f"{PAGE}/template").content)
    assert response.status_code == 400
    assert response.json() == {"detail": "El archivo no contiene proveedores"}
    assert supplier_count() == before


def test_pagina_de_carga(admin):
    response = admin.get(PAGE)
    assert response.status_code == 200
    for fragment in (
        "Algunas filas contienen errores o no se han podido leer correctamente. ¿Desea agregar las filas válidas?",
        "No, corregir primero (Recomendado)",
        "Agrega las filas válidas y omite el resto",
        'accept=".xlsx"',
        "/static/js/supplier_import.js",
        "<noscript>",
    ):
        assert fragment in response.text
    assert 'href="/suppliers/import"' in admin.get("/suppliers").text


@pytest.mark.parametrize("email", ["pmo@poc.local", "proveedor1@poc.local"])
def test_solo_el_administrador_accede(client, email):
    login(client, email)
    before = supplier_count()
    assert client.get(PAGE).status_code == 403
    assert client.get(f"{PAGE}/template").status_code == 403
    response = client.post(
        PAGE,
        data={"csrf_token": csrf(client, "/")},
        files={"upload": ("proveedores.xlsx", workbook([national()]), "application/octet-stream")},
        headers={"Accept": "application/json"},
    )
    assert response.status_code == 403
    assert supplier_count() == before


def test_envio_sin_token_csrf(admin):
    before = supplier_count()
    response = post_import(admin, workbook([national()]), token=False)
    assert response.status_code == 403
    assert supplier_count() == before


# --- Validacion del archivo --------------------------------------------------------------------------------------


@pytest.mark.parametrize("filename", ["proveedores.csv", "proveedores.xls", "proveedores.xlsm", "sin_extension"])
def test_extension_no_permitida(admin, filename):
    response = post_import(admin, workbook([national()]), filename=filename)
    assert response.status_code == 400
    assert response.json()["detail"] == "Solo se aceptan archivos de Excel (.xlsx) generados con la plantilla"


def zip_with(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in members.items():
            package.writestr(name, data)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "content",
    [
        b"%PDF-1.7\nno es un libro de Excel",
        zip_with({"documento.txt": b"texto"}),
        zip_with({"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<workbook>roto"}),
    ],
    ids=["pdf", "zip-sin-libro", "libro-danado"],
)
def test_contenido_que_no_es_un_libro(admin, content):
    response = post_import(admin, content)
    assert response.status_code == 400
    assert response.json()["detail"] == "El archivo no es un libro de Excel válido"


@pytest.mark.parametrize(
    ("headers", "sheet"),
    [
        (HEADERS[:-1], "Proveedores"),
        ((HEADERS[1], HEADERS[0], *HEADERS[2:]), "Proveedores"),
        (("Tipo", *HEADERS[1:]), "Proveedores"),
        ((*HEADERS, "Extra"), "Proveedores"),
        (HEADERS, "Hoja1"),
    ],
    ids=["falta-columna", "orden", "nombre", "columna-adicional", "sin-hoja"],
)
def test_plantilla_distinta(admin, headers, sheet):
    response = post_import(admin, workbook([national()], headers=headers, sheet=sheet))
    assert response.status_code == 400
    assert response.json()["detail"] == MSG_TEMPLATE


def test_encabezados_sin_distinguir_mayusculas_ni_espacios(admin):
    headers = [f"  {header.upper()} " for header in HEADERS]
    row = national()
    content = workbook(
        [{header: row.get(original) for header, original in zip(headers, HEADERS, strict=True)}], headers
    )
    response = post_import(admin, content)
    assert response.status_code == 200
    assert response.json()["created"] == 1


def test_demasiadas_filas(admin):
    before = supplier_count()
    response = post_import(admin, workbook([national() for _ in range(1001)]))
    assert response.status_code == 400
    assert response.json() == {"detail": "El archivo excede el máximo de 1000 proveedores por carga"}
    assert supplier_count() == before


def test_archivo_demasiado_grande(admin):
    response = post_import(admin, b"x" * (5 * 1024 * 1024 + 1))
    assert response.status_code == 400
    assert response.json()["detail"] == "El archivo excede 5 MB"


def test_bomba_de_descompresion(admin):
    content = zip_with(
        {"[Content_Types].xml": b"<Types/>", "xl/workbook.xml": b"<workbook/>", "xl/relleno.bin": bytes(51 << 20)}
    )
    assert len(content) < 5 * 1024 * 1024
    response = post_import(admin, content)
    assert response.status_code == 400
    assert response.json()["detail"] == "El contenido del archivo excede el tamaño permitido"


def test_entidades_xml(admin):
    laughs = b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;">]>'
    source = zipfile.ZipFile(io.BytesIO(workbook([national()])))
    members = {}
    for name in source.namelist():
        data = source.read(name)
        if name == "xl/worksheets/sheet1.xml":
            data = laughs + (data.split(b"?>", 1)[1] if data.startswith(b"<?xml") else data)
        members[name] = data
    response = post_import(admin, zip_with(members))
    assert response.status_code == 400
    assert response.json()["detail"] == "El archivo no es un libro de Excel válido"


# --- Mapeo y normalizacion ---------------------------------------------------------------------------------------


def test_fila_nacional_normalizada(admin):
    rfc, email = unique_rfc(), f"Contacto-{uuid4().hex[:8]}@SDN.mx"
    row = national(**{"Razón social": "  Servicios   Digitales del Norte SA de CV ", "RFC": rfc.lower()})
    row["Correo electrónico"] = email
    assert post_import(admin, workbook([row])).status_code == 200
    supplier = supplier_by(rfc=rfc)
    assert supplier.origin == SupplierOrigin.NATIONAL
    assert supplier.supplier_type == SupplierType.PERSONA_MORAL
    assert supplier.business_name == "Servicios Digitales del Norte SA de CV"
    assert (supplier.country, supplier.email, supplier.foreign_tax_id) == ("MX", email.lower(), None)


def test_fila_internacional_normalizada(admin):
    tax_id = f"12-{uuid4().hex[:7]}"
    row = international(**{"Origen": "INTERNACIONAL", "Tipo de persona": "moral", "País": "us"})
    row["Identificador fiscal extranjero"] = tax_id
    assert post_import(admin, workbook([row])).status_code == 200
    supplier = supplier_by(foreign_tax_id=tax_id.upper())
    assert supplier.origin == SupplierOrigin.INTERNATIONAL
    assert (supplier.rfc, supplier.country, supplier.supplier_type) == (None, "US", SupplierType.PERSONA_MORAL)
    assert supplier.tax_identifier == f"US {tax_id.upper()}"


def test_valores_de_lista_sin_acentos_ni_mayusculas(admin):
    rfc = unique_rfc(moral=False)
    row = national(**{"Tipo de persona": "FISICA", "RFC": rfc, "Convenio de confidencialidad": "si"})
    assert post_import(admin, workbook([row])).status_code == 200
    supplier = supplier_by(rfc=rfc)
    assert supplier.supplier_type == SupplierType.PERSONA_FISICA
    assert supplier.confidentiality_agreement is True


def test_telefono_capturado_como_numero(admin):
    row = national(**{"Teléfono": 5512345678, "Notas": "Alta desde la carga masiva"})
    assert post_import(admin, workbook([row])).status_code == 200
    supplier = supplier_by(rfc=row["RFC"])
    assert (supplier.phone, supplier.notes, supplier.confidentiality_agreement) == (
        "5512345678",
        "Alta desde la carga masiva",
        False,
    )


# --- Validacion de cada fila -------------------------------------------------------------------------------------


def test_errores_de_varias_filas_en_una_sola_respuesta(admin):
    before = supplier_count()
    rows = [national(**{"Correo electrónico": None}), national(), national(RFC="ABC123")]
    response = post_import(admin, workbook(rows))
    assert response.status_code == 400
    assert texts(response) == ["Fila 2 · Correo electrónico: es obligatorio", "Fila 4 · RFC: tiene un formato inválido"]
    assert (response.json()["registrable"], response.json()["invalid"]) == (1, 2)
    assert supplier_count() == before


@pytest.mark.parametrize(
    ("row", "expected"),
    [
        (national(RFC="GOMA850612H45"), "RFC: una persona moral debe tener un RFC de 12 caracteres"),
        (
            national(**{"Tipo de persona": "Física", "RFC": "GOM850612H45"}),
            "RFC: una persona física debe tener un RFC de 13 caracteres",
        ),
        (
            national(**{"Tipo de persona": "Física", "RFC": "XEXX010101000"}),
            "RFC: no se permite un RFC genérico; un proveedor extranjero se registra con Origen Internacional",
        ),
        (national(RFC="SDN201345AB1"), "RFC: tiene un formato inválido"),
        (national(RFC=None), "RFC: es obligatorio para proveedores nacionales"),
        (
            national(**{"Identificador fiscal extranjero": "X-1"}),
            "Identificador fiscal extranjero: debe quedar vacío para proveedores nacionales",
        ),
        (national(**{"País": "US"}), "País: debe quedar vacío o ser MX para proveedores nacionales"),
        (international(**{"País": "USA"}), "País: use el código ISO de dos letras (p. ej. US)"),
        (international(**{"País": "MX"}), "País: un proveedor internacional no puede tener país MX"),
        (international(**{"País": None}), "País: es obligatorio para proveedores internacionales"),
        (
            international(**{"Identificador fiscal extranjero": "ID#123"}),
            "Identificador fiscal extranjero: use hasta 40 caracteres: letras, dígitos, espacios, puntos, guiones o "
            "diagonales",
        ),
        (national(Origen=None), "Origen: es obligatorio"),
        (national(Origen="Local"), "Origen: use Nacional o Internacional"),
        (national(**{"Tipo de persona": "Otra"}), "Tipo de persona: use Física o Moral"),
        (national(**{"Razón social": "X"}), "Razón social: es demasiado corto"),
        (national(**{"Razón social": "X" * 251}), "Razón social: es demasiado largo"),
        (national(**{"Razón social": None}), "Razón social: es obligatorio"),
        (national(**{"Correo electrónico": "no-es-correo"}), "Correo electrónico: no es un correo válido"),
        (
            national(**{"Correo electrónico": f"{'a' * 60}@{'b' * 60}.{'c' * 60}.{'d' * 60}.{'e' * 60}.mx"}),
            "Correo electrónico: es demasiado largo",
        ),
        (national(**{"Teléfono": "55-12"}), "Teléfono: use de 7 a 30 caracteres: dígitos, espacios, +, (, ) o -"),
        (national(Notas="n" * 1001), "Notas: admite a lo sumo 1000 caracteres"),
        (national(**{"Convenio de confidencialidad": "Tal vez"}), "Convenio de confidencialidad: use Sí o No"),
        (
            national(**{"Correo electrónico": '=A2&"@proveedor.mx"'}),
            "Correo electrónico: la celda contiene una fórmula; capture el valor",
        ),
        (
            national(**{"Teléfono": date(2026, 1, 1)}),
            "Teléfono: la celda contiene una fecha; capture el valor como texto",
        ),
    ],
)
def test_reglas_por_fila(admin, row, expected):
    response = post_import(admin, workbook([row]))
    assert response.status_code == 400
    assert texts(response) == [f"Fila 2 · {expected}"]


def test_proveedor_internacional_incompleto(admin):
    row = international(**{"Identificador fiscal extranjero": None, "RFC": unique_rfc()})
    response = post_import(admin, workbook([row]))
    assert texts(response) == [
        "Fila 2 · RFC: debe quedar vacío para proveedores internacionales",
        "Fila 2 · Identificador fiscal extranjero: es obligatorio para proveedores internacionales",
    ]


# --- Duplicados --------------------------------------------------------------------------------------------------


def test_rfc_repetido_en_el_archivo(admin):
    repeated = unique_rfc()
    rows = [national(), national(RFC=repeated), national(), national(), national(), national(RFC=repeated)]
    content = workbook(rows)
    response = post_import(admin, content)
    assert response.status_code == 400
    assert texts(response) == [
        "Fila 3 · RFC: repetido en las filas 3 y 7",
        "Fila 7 · RFC: repetido en las filas 3 y 7",
    ]
    partial = post_import(admin, content, mode="partial", sha256=response.json()["sha256"])
    assert partial.json()["created"] == 4
    assert supplier_by(rfc=repeated) is None


def test_identificador_y_correo_repetidos_en_el_archivo(admin):
    first = international()
    second = international(**{"Identificador fiscal extranjero": first["Identificador fiscal extranjero"]})
    email = unique_email()
    rows = [first, second, national(**{"Correo electrónico": email}), national(**{"Correo electrónico": email.upper()})]
    rows += [national() for _ in range(4)]
    response = post_import(admin, workbook(rows))
    assert texts(response) == [
        "Fila 2 · Identificador fiscal extranjero: repetido para el mismo país en las filas 2 y 3",
        "Fila 3 · Identificador fiscal extranjero: repetido para el mismo país en las filas 2 y 3",
        "Fila 4 · Correo electrónico: repetido en las filas 4 y 5",
        "Fila 5 · Correo electrónico: repetido en las filas 4 y 5",
    ]
    assert service._join_numbers([2, 3, 4, 5, 6, 7, 8]) == "2, 3, 4, 5, 6 y 2 más"


def test_proveedor_existente_omitido(admin):
    existing = existing_supplier()
    rows = [
        national(RFC=existing.rfc, **{"Razón social": "Otro Nombre SA de CV", "Correo electrónico": unique_email()}),
        national(),
        national(),
    ]
    response = post_import(admin, workbook(rows))
    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 2
    assert body["skipped"] == [{"row": 2, "identifier": existing.rfc, "business_name": "Otro Nombre SA de CV"}]
    unchanged = supplier_by(id=existing.id)
    assert (unchanged.business_name, unchanged.email) == (existing.business_name, existing.email)


def test_proveedor_internacional_existente_omitido(admin):
    tax_id = f"EX-{uuid4().hex[:8].upper()}"
    existing_supplier(
        rfc=None, origin=SupplierOrigin.INTERNATIONAL, foreign_tax_id=tax_id, country="CA", email=unique_email()
    )
    rows = [international(**{"Identificador fiscal extranjero": tax_id, "País": "CA"}), international()]
    body = post_import(admin, workbook(rows)).json()
    assert (body["created"], body["skipped"][0]["identifier"]) == (1, f"CA {tax_id}")


@pytest.mark.parametrize("owner", ["supplier", "user"])
def test_correo_en_uso(admin, owner):
    email = unique_email()
    if owner == "supplier":
        existing_supplier(email=email)
    else:
        with SessionLocal() as db:
            db.add(User(name="Usuario existente", email=email, password_hash="x", role="PMO"))
            db.commit()
    before = supplier_count()
    response = post_import(admin, workbook([national(**{"Correo electrónico": email.upper()})]))
    assert response.status_code == 400
    assert texts(response) == ["Fila 2 · Correo electrónico: ya está registrado para otro proveedor o usuario"]
    assert supplier_count() == before


# --- Confirmacion, registro atomico y resumen --------------------------------------------------------------------


def test_confirmacion_y_registro_de_las_filas_validas(admin):
    valid = [national() for _ in range(8)]
    content = workbook([*valid[:4], national(RFC="MAL"), *valid[4:], national(**{"Correo electrónico": "x"})])
    before = supplier_count()
    strict = post_import(admin, content)
    body = strict.json()
    assert strict.status_code == 400
    assert (body["registrable"], body["invalid"], body["rows"]) == (8, 2, 10)
    assert body["sha256"] == hashlib.sha256(content).hexdigest()
    assert body["detail"] == "Algunas filas contienen errores o no se han podido leer correctamente."
    assert supplier_count() == before
    partial = post_import(admin, content, mode="partial", sha256=body["sha256"])
    assert partial.status_code == 200
    result = partial.json()
    assert (result["result"], result["mode"], result["created"], result["invalid"]) == ("partial", "partial", 8, 2)
    assert texts(partial) == [
        "Fila 6 · RFC: tiene un formato inválido",
        "Fila 11 · Correo electrónico: no es un correo válido",
    ]
    assert supplier_count() == before + 8


@pytest.mark.parametrize("sha256", [None, "0" * 64])
def test_archivo_distinto_al_validado(admin, sha256):
    before = supplier_count()
    response = post_import(admin, workbook([national(), national(RFC="MAL")]), mode="partial", sha256=sha256)
    assert response.status_code == 409
    assert response.json() == {"detail": "El archivo cambió desde la validación. Vuelva a cargarlo."}
    assert supplier_count() == before


def test_ninguna_fila_registrable(admin):
    existing = existing_supplier()
    response = post_import(admin, workbook([national(RFC="MAL"), national(RFC=existing.rfc)]))
    assert response.status_code == 400
    assert response.json()["registrable"] == 0


def test_mas_de_200_errores(admin):
    rows = [national(**{"Correo electrónico": "no-es-correo"}) for _ in range(350)]
    body = post_import(admin, workbook(rows)).json()
    assert (len(body["errors"]), body["hidden_errors"]) == (200, 150)
    assert [error["row"] for error in body["errors"]] == list(range(2, 202))


def test_modo_invalido_y_archivo_ausente(admin):
    assert post_import(admin, workbook([national()]), mode="otro").json() == {"detail": "Modo de carga no válido"}
    response = admin.post(PAGE, data={"csrf_token": csrf(admin, PAGE)}, headers={"Accept": "application/json"})
    assert response.status_code == 400
    assert response.json()["detail"] == "Solo se aceptan archivos de Excel (.xlsx) generados con la plantilla"


def test_archivo_sin_errores_se_registra_en_un_paso(admin):
    response = post_import(admin, workbook([national(), None, international(), national()]))
    assert response.status_code == 200
    body = response.json()
    assert (body["result"], body["created"], body["rows"]) == ("imported", 3, 3)


def test_conflicto_de_unicidad_revierte_todo(admin, monkeypatch):
    existing = existing_supplier()
    new = national()
    monkeypatch.setattr(service, "classify", lambda db, rows: None)  # simula otro proceso que ya registro el RFC
    response = post_import(admin, workbook([new, national(RFC=existing.rfc)]))
    assert response.status_code == 409
    assert response.json() == {
        "detail": "Otro proceso registró proveedores de este archivo mientras se procesaba. Vuelva a cargarlo."
    }
    assert supplier_by(rfc=new["RFC"]) is None


def test_proveedor_recien_cargado(admin):
    row = national()
    assert post_import(admin, workbook([row])).status_code == 200
    supplier = supplier_by(rfc=row["RFC"])
    assert supplier.status == SupplierStatus.REGISTERED
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(User).where(User.supplier_id == supplier.id)) == 0
    assert "Registrado" in admin.get("/suppliers").text
    detail = admin.get(f"/suppliers/{supplier.id}").text
    assert "Registrado" in detail and row["RFC"] in detail and "Nacional" in detail


def mixed_file(errors: int = 0) -> tuple[bytes, list[Supplier]]:
    existing = [existing_supplier(), existing_supplier()]
    rows = [national() for _ in range(8 - errors)]
    rows[1:1] = [national(RFC=supplier.rfc) for supplier in existing]
    rows += [national(RFC=f"MAL{index}") for index in range(errors)]
    return workbook(rows), existing


def test_resumen_de_proveedores_nuevos_y_existentes(admin):
    content, existing = mixed_file()
    body = post_import(admin, content).json()
    assert body["summary"] == "10 filas leídas · 8 proveedores registrados · 2 omitidos"
    assert [entry["identifier"] for entry in body["skipped"]] == [supplier.rfc for supplier in existing]
    again = post_import(admin, content).json()
    assert again["summary"] == "10 filas leídas · 0 proveedores registrados · 10 omitidos"


def test_resumen_de_una_carga_parcial(admin):
    content, _ = mixed_file(errors=2)
    sha256 = post_import(admin, content).json()["sha256"]
    body = post_import(admin, content, mode="partial", sha256=sha256).json()
    assert body["summary"] == "10 filas leídas · 6 proveedores registrados · 2 omitidos · 2 con errores"
    assert len(body["skipped"]) == 2 and len(body["errors"]) == 2


def test_resumen_en_singular():
    assert service.summary_text(1, 1, 1, 0) == "1 fila leída · 1 proveedor registrado · 1 omitido"


# --- Auditoria y log ---------------------------------------------------------------------------------------------


def bulk_audit(sha256: str) -> AuditLog:
    with SessionLocal() as db:
        return db.scalar(
            select(AuditLog).where(AuditLog.action == "SUPPLIER_BULK_IMPORTED", AuditLog.entity_id == sha256)
        )


def created_audits(sha256: str) -> int:
    with SessionLocal() as db:
        return db.scalar(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "SUPPLIER_CREATED", AuditLog.new_value["import_sha256"].astext == sha256)
        )


def admin_id() -> int:
    with SessionLocal() as db:
        return db.scalar(select(User.id).where(User.email == "admin@poc.local"))


def test_auditoria_de_una_carga_concluida(admin):
    content, _ = mixed_file()
    sha256 = hashlib.sha256(content).hexdigest()
    assert post_import(admin, content).status_code == 200
    entry = bulk_audit(sha256)
    assert entry.user_id == admin_id()
    assert entry.new_value == {
        "mode": "strict",
        "sha256": sha256,
        "size_bytes": len(content),
        "rows": 10,
        "created": 8,
        "skipped": 2,
        "invalid": 0,
    }
    assert created_audits(sha256) == 8


def test_auditoria_de_una_carga_parcial(admin):
    content, _ = mixed_file(errors=2)
    sha256 = post_import(admin, content).json()["sha256"]
    post_import(admin, content, mode="partial", sha256=sha256)
    entry = bulk_audit(sha256)
    assert (entry.new_value["mode"], entry.new_value["created"], entry.new_value["skipped"]) == ("partial", 6, 2)
    assert entry.new_value["invalid"] == 2
    assert created_audits(sha256) == 6


def test_carga_rechazada_sin_auditoria(admin):
    before = audit_count()
    post_import(admin, workbook([national(RFC="MAL"), national()]))
    post_import(admin, b"no es excel")
    assert audit_count() == before


def test_eventos_de_carga_en_el_log(admin):
    offset = log_offset()
    post_import(admin, workbook([national(), national(), international()]))
    content = workbook([national(), national(RFC="MAL"), national(RFC="MAL2")])
    sha256 = post_import(admin, content).json()["sha256"]
    post_import(admin, content, mode="partial", sha256=sha256)
    events = [event for event in events_since(offset) if event.get("event") == "supplier.bulk_import"]
    imported, rejected, partial = events
    assert {key: imported[key] for key in ("mode", "result", "rows", "registered", "skipped", "invalid")} == {
        "mode": "strict",
        "result": "imported",
        "rows": 3,
        "registered": 3,
        "skipped": 0,
        "invalid": 0,
    }
    assert imported["size_bytes"] > 0 and imported["duration_ms"] >= 0
    assert (rejected["result"], rejected["invalid"]) == ("rejected", 2)
    assert (partial["mode"], partial["result"], partial["invalid"], partial["registered"]) == (
        "partial",
        "partial",
        2,
        1,
    )


def test_log_sin_datos_del_archivo(admin):
    offset = log_offset()
    row, foreign = national(**{"Razón social": "Razon Social Confidencial SA"}), international()
    filename = "catalogo_confidencial_proveedores.xlsx"
    post_import(admin, workbook([row, foreign]), filename=filename)
    post_import(admin, workbook([national(RFC="MAL"), row]), filename=filename)
    post_import(admin, b"<!DOCTYPE no es excel", filename=filename)
    content = json.dumps(events_since(offset), ensure_ascii=False)
    assert "supplier.bulk_import" in content
    forbidden = [
        row["RFC"],
        row["Correo electrónico"],
        "Razon Social Confidencial SA",
        foreign["Identificador fiscal extranjero"],
        foreign["Correo electrónico"],
        filename,
    ]
    for value in forbidden:
        assert value not in content, value
