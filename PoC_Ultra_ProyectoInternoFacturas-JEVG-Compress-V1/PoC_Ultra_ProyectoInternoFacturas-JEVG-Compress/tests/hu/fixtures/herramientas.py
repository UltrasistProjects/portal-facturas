"""Herramientas de apoyo de las pruebas funcionales de las HUs (Playwright). Se ejecutan con el Python del proyecto
(.venv), porque reutilizan openpyxl, PyMuPDF y la configuracion de la aplicacion.

Subcomandos:
  fixtures <run> <dir>                       genera los archivos de prueba de la ejecucion (XML, PDF, TXT)
  plantilla-proveedores <xlsx> <out> <run> <modo>   llena la plantilla descargada del portal (modo: validos|errores)
  plantilla-catalogo <xlsx> <out> <clave> <descripcion>   agrega una fila a la plantilla de un catalogo
  inspeccionar-xlsx <xlsx>                   hojas, encabezados y filas de datos de un libro
  correo <outbox> <destinatario> <texto-asunto> [desde-iso]   ultimo .eml para ese destinatario, como JSON
  usuario <correo>                           datos de credenciales del usuario en la BD del portal, como JSON

Ningun subcomando modifica la base de datos.
"""

import base64
import json
import re
import shutil
import sys
import uuid
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from email.utils import getaddresses
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
DEMO = ROOT / "data" / "demo_documents"


def _pdf(path: Path, lines: list[str]) -> None:
    import fitz  # PyMuPDF

    document = fitz.open()
    page = document.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=11)
        y += 18
    document.save(path)
    document.close()


def _cfdi(uuid_value: str, postal_code: str = "03930") -> str:
    raw = (DEMO / "cfdi_demo_correcto.xml").read_text(encoding="utf-8")
    raw = re.sub(r'UUID="[^"]*"', f'UUID="{uuid_value}"', raw)
    return raw.replace('DomicilioFiscalReceptor="03930"', f'DomicilioFiscalReceptor="{postal_code}"')


def _payment_complement(related_uuid: str) -> str:
    """CFDI de pago (tipo P) que liquida la factura `related_uuid` (HU-23), con un UUID propio."""
    raw = (DEMO / "complemento_pago_demo.xml").read_text(encoding="utf-8")
    raw = raw.replace("DEMO0001-0000-4000-8000-000000000001", related_uuid)
    return re.sub(r'UUID="DEMOPAGO-[^"]*"', f'UUID="{str(uuid.uuid4()).upper()}"', raw)


def fixtures(run: str, out: str) -> None:
    folder = Path(out)
    folder.mkdir(parents=True, exist_ok=True)
    uuid_ok = str(uuid.uuid4()).upper()
    files = {
        # Nacional: CFDI correcto (UUID nuevo), CFDI con el codigo postal del receptor distinto (XML-010) y un CFDI con
        # el mismo UUID que el correcto (duplicado, RN-HU13-01).
        "cfdi_ok": folder / f"cfdi_{run}_correcto.xml",
        "cfdi_cp": folder / f"cfdi_{run}_cp_06600.xml",
        "cfdi_dup": folder / f"cfdi_{run}_uuid_repetido.xml",
        "cfdi_ok_b": folder / f"cfdi_{run}_correcto_b.xml",
        "cfdi_ok_c": folder / f"cfdi_{run}_correcto_c.xml",
        # HU-23: CFDI nacional con MetodoPago PPD (el del demo) y su Complemento de Pago (tipo P) que lo relaciona.
        "cfdi_ppd": folder / f"cfdi_{run}_ppd.xml",
        "complemento_pago": folder / f"complemento_pago_{run}.xml",
        "pdf_cfdi": folder / f"factura_{run}.pdf",
        "orden_compra": folder / f"orden_compra_{run}.txt",
        "vobo": folder / f"vobo_{run}.txt",
        "invoice": folder / f"INV-QA-{run}.pdf",
        "invoice_otro": folder / f"INV-QA-{run}-B.pdf",
        "acuse": folder / f"acuse_cancelacion_{run}.pdf",
        "requisito_alta": folder / f"requisito_alta_{run}.pdf",
        "texto_invalido": folder / f"no_es_xml_{run}.txt",
    }
    files["cfdi_ok"].write_text(_cfdi(uuid_ok), encoding="utf-8")
    files["cfdi_cp"].write_text(_cfdi(str(uuid.uuid4()).upper(), "06600"), encoding="utf-8")
    files["cfdi_dup"].write_text(_cfdi(uuid_ok), encoding="utf-8")
    files["cfdi_ok_b"].write_text(_cfdi(str(uuid.uuid4()).upper()), encoding="utf-8")
    files["cfdi_ok_c"].write_text(_cfdi(str(uuid.uuid4()).upper()), encoding="utf-8")
    uuid_ppd = str(uuid.uuid4()).upper()
    files["cfdi_ppd"].write_text(_cfdi(uuid_ppd), encoding="utf-8")
    files["complemento_pago"].write_text(_payment_complement(uuid_ppd), encoding="utf-8")
    shutil.copyfile(DEMO / "factura_demo.pdf", files["pdf_cfdi"])
    shutil.copyfile(DEMO / "orden_compra_demo.txt", files["orden_compra"])
    shutil.copyfile(DEMO / "vobo_demo.txt", files["vobo"])
    # Invoice internacional: trae el identificador fiscal del proveedor demo y el codigo postal de ULTRASIST, pero no
    # su razon social, para que INT-002 resulte advertencia (no bloquea el envio).
    invoice_lines = [
        f"INVOICE INV-QA-{run}",
        "Global Data Services Inc. - Tax ID 98-7654321",
        "Bill to: customer in Mexico City, ZIP 03930",
        "Data analytics services - August 2026",
        "Subtotal 1,000.00 USD  Tax 0.00  Total 1,000.00 USD",
        "QA TEST DOCUMENT - NO FISCAL VALIDITY",
    ]
    _pdf(files["invoice"], invoice_lines)
    _pdf(files["invoice_otro"], invoice_lines)
    _pdf(
        files["acuse"],
        [f"ACUSE DE CANCELACION (PRUEBA QA {run})", "Documento de prueba sin validez fiscal."],
    )
    _pdf(
        files["requisito_alta"],
        [f"DOCUMENTO DEL EXPEDIENTE (PRUEBA QA {run})", "Requisito de alta del proveedor (HU-21). Sin validez legal."],
    )
    files["texto_invalido"].write_text("Este archivo no es un XML.\n", encoding="utf-8")
    print(json.dumps({key: str(path) for key, path in files.items()} | {"uuid_ok": uuid_ok, "uuid_ppd": uuid_ppd}))


def _rfc(prefix: str, run: str, physical: bool) -> str:
    # Patron SAT: 3 letras (moral) o 4 (fisica), fecha AAMMDD valida y 3 caracteres de homoclave unicos por ejecucion.
    letters = (prefix * 2)[: 4 if physical else 3]
    return f"{letters}260101{run[-3:]}"


def plantilla_proveedores(source: str, out: str, run: str, mode: str) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(source)
    sheet = workbook["Proveedores"]
    domain = "proveedores-qa.example"
    if mode == "validos":
        rows = [
            (
                "Nacional",
                "Moral",
                f"QA Alfa Servicios {run} SA de CV",
                _rfc("QAA", run, False),
                None,
                None,
                f"alfa.{run}@{domain}".lower(),
                "55 1234 5678",
                "Sí",
                "Alta por prueba funcional HU-01",
            ),
            (
                "Nacional",
                "Física",
                f"QA Bravo Persona {run}",
                _rfc("QABR", run, True),
                None,
                None,
                f"bravo.{run}@{domain}".lower(),
                None,
                "No",
                None,
            ),
            (
                "Internacional",
                "Moral",
                f"QA Charlie Consulting {run} LLC",
                None,
                f"QA-{run}",
                "US",
                f"charlie.{run}@{domain}".lower(),
                "+1 206 555 0100",
                "Sí",
                "Proveedor internacional de prueba",
            ),
        ]
    else:
        rows = [
            (
                "Nacional",
                "Moral",
                f"QA Delta Valida {run} SA de CV",
                _rfc("QAD", run, False),
                None,
                None,
                f"delta.{run}@{domain}".lower(),
                None,
                "Sí",
                "Fila valida",
            ),
            (
                "Nacional",
                "Moral",
                f"QA Echo RFC Invalido {run}",
                "ABC123",
                None,
                None,
                f"echo.{run}@{domain}".lower(),
                None,
                None,
                "Error: RFC con formato invalido",
            ),
            (
                "Internacional",
                "Moral",
                f"QA Foxtrot Sin ID {run}",
                None,
                None,
                "US",
                f"foxtrot.{run}@{domain}".lower(),
                None,
                None,
                "Error: falta el identificador fiscal",
            ),
            (
                "Nacional",
                "Moral",
                f"QA Golf Correo {run}",
                _rfc("QAG", run, False),
                None,
                None,
                "correo-sin-arroba.example",
                None,
                None,
                "Error: correo invalido",
            ),
        ]
    for row in rows:
        sheet.append(row)
    workbook.save(out)
    print(json.dumps({"filas": [{"razon_social": r[2], "rfc": r[3], "id_fiscal": r[4], "correo": r[6]} for r in rows]}))


def inspeccionar_xlsx(source: str) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(source)
    first = workbook[workbook.sheetnames[0]]
    headers = [cell.value for cell in first[1]]
    print(json.dumps({"hojas": workbook.sheetnames, "encabezados": headers, "filas_datos": first.max_row - 1}))


def plantilla_catalogo(source: str, out: str, code: str, description: str) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(source)
    sheet = workbook["Catalogo"]
    sheet.append((code, description, "Sí"))
    workbook.save(out)
    print(json.dumps({"filas": sheet.max_row - 1}))


def _html_with_inline_images(message) -> str | None:
    html_part = message.get_body(preferencelist=("html",))
    if html_part is None:
        return None
    html = html_part.get_content()
    for part in message.walk():
        cid = part.get("Content-ID")
        if cid and part.get_content_maintype() == "image":
            data = base64.b64encode(part.get_payload(decode=True)).decode()
            html = html.replace(f"cid:{cid.strip('<>')}", f"data:{part.get_content_type()};base64,{data}")
    return html


def correo(outbox: str, recipient: str, subject_text: str, since: str | None = None) -> None:
    """Ultimo correo (por nombre de archivo, que empieza con la fecha UTC) para el destinatario cuyo asunto contiene
    el texto. `since` (ISO UTC) descarta los anteriores a esa fecha."""
    folder = Path(outbox)
    limit = datetime.fromisoformat(since) if since else None
    for path in sorted(folder.glob("*.eml"), reverse=True):
        stamp = datetime.strptime(path.name[:21], "%Y%m%dT%H%M%S%f").replace(tzinfo=timezone.utc)
        if limit and stamp < limit:
            break
        with path.open("rb") as handle:
            message = BytesParser(policy=policy.default).parse(handle)
        to = [address.lower() for _, address in getaddresses(message.get_all("To", []))]
        cc = [address.lower() for _, address in getaddresses(message.get_all("Cc", []))]
        if recipient.lower() not in to + cc or subject_text.lower() not in str(message["Subject"]).lower():
            continue
        text_part = message.get_body(preferencelist=("plain",))
        print(
            json.dumps(
                {
                    "archivo": str(path),
                    "fecha_utc": stamp.isoformat(),
                    "de": str(message.get("From", "")),
                    "para": to,
                    "cc": cc,
                    "asunto": str(message["Subject"]),
                    "texto": text_part.get_content() if text_part else "",
                    "html": _html_with_inline_images(message),
                },
                ensure_ascii=False,
            )
        )
        return
    print(json.dumps(None))


def usuario(email: str) -> None:
    sys.path.insert(0, str(ROOT))
    from sqlalchemy import select

    from app.core.database import SessionLocal
    from app.models import User

    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == email.lower()))
        if user is None:
            print(json.dumps(None))
            return
        print(
            json.dumps(
                {
                    "correo": user.email,
                    "rol": user.role.value,
                    "password_hash_vacio": not user.password_hash,
                    "keycloak_sub_registrado": bool(user.keycloak_sub),
                    "proveedor_id": user.supplier_id,
                }
            )
        )


if __name__ == "__main__":
    command, *args = sys.argv[1:]
    {
        "fixtures": fixtures,
        "plantilla-proveedores": plantilla_proveedores,
        "plantilla-catalogo": plantilla_catalogo,
        "inspeccionar-xlsx": inspeccionar_xlsx,
        "correo": correo,
        "usuario": usuario,
    }[command](*args)
